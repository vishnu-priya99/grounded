"""Format-specific parsers. Each returns a :class:`DocModel`.

Torch-free by design: PyMuPDF, python-docx, python-pptx and RapidOCR (ONNX) do
all the heavy lifting so the whole stack runs comfortably on a 16 GB laptop with
no GPU. Docling would give better layout fidelity and is the documented upgrade
path.
"""

from __future__ import annotations

import functools
import re
import statistics
from pathlib import Path

from grounded.config import settings
from grounded.ingest.models import DocModel, Element, SourceKind
from grounded.llm import vision_available
from grounded.util import clean_text, stable_id

# ---------------------------------------------------------------------------
# extension routing
# ---------------------------------------------------------------------------

DOC_EXT = {".pdf", ".docx", ".txt", ".md", ".markdown", ".pptx", ".html", ".htm"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"}
CODE_EXT = {
    ".py": "python", ".js": "javascript", ".ts": "typescript", ".tsx": "tsx",
    ".java": "java", ".go": "go", ".rs": "rust", ".c": "c", ".h": "c",
    ".cpp": "cpp", ".cc": "cpp", ".cs": "csharp", ".rb": "ruby", ".php": "php",
    ".sh": "bash", ".sql": "sql", ".scala": "scala", ".kt": "kotlin",
    ".swift": "swift", ".r": "r", ".m": "matlab", ".lua": "lua",
}
SPREADSHEET_EXT = {".csv", ".xlsx", ".xls", ".tsv"}
_TEXTLIKE_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".text"}

# magic-byte signatures, so a mislabeled upload still routes correctly
_MAGIC: list[tuple[bytes, str]] = [
    (b"%PDF-", "pdf"),
    (b"\x89PNG\r\n\x1a\n", "image"),
    (b"\xff\xd8\xff", "image"),
    (b"GIF8", "image"),
    (b"BM", "image"),
    (b"PK\x03\x04", "zip"),  # docx / xlsx / pptx / plain zip
]


def supported(path: Path) -> bool:
    ext = path.suffix.lower()
    if (
        ext in DOC_EXT or ext in IMAGE_EXT or ext in CODE_EXT
        or ext in SPREADSHEET_EXT or ext in _TEXTLIKE_EXT
    ):
        return True
    return _sniff(path) is not None


def _sniff(path: Path) -> str | None:
    """Best-effort content type from the first bytes. Returns one of
    ``pdf | image | zip | text`` or ``None``."""
    try:
        head = path.read_bytes()[:2048]
    except OSError:
        return None
    if not head:
        return None
    for sig, kind in _MAGIC:
        if head.startswith(sig):
            return kind
    # mostly-printable -> treat as text
    printable = sum(1 for b in head if b in (9, 10, 13) or 32 <= b < 127)
    return "text" if printable / len(head) > 0.85 else None


def _resolve_kind(path: Path) -> str:
    """Decide how to parse a file: extension first, content sniff as tiebreak
    or override for the unambiguous binary signatures."""
    ext = path.suffix.lower()
    sniff = _sniff(path)
    if sniff == "pdf":
        return "pdf"
    if sniff == "image" and ext not in CODE_EXT:
        return "image"
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext == ".pptx":
        return "pptx"
    if ext in {".html", ".htm"}:
        return "html"
    if ext in _TEXTLIKE_EXT:
        return "text"
    if ext in IMAGE_EXT:
        return "image"
    if ext in CODE_EXT:
        return "code"
    if sniff == "zip":  # PK header but unknown office ext
        return "docx"
    if sniff == "text":
        return "text"
    raise ValueError(f"unrecognised file type: {path.name}")


def guard(path: Path) -> None:
    if not path.exists():
        raise ValueError("file not found")
    size = path.stat().st_size
    if size == 0:
        raise ValueError("file is empty")
    if size > settings.max_file_mb * 1024 * 1024:
        raise ValueError(
            f"file is {size / 1e6:.0f} MB, over the {settings.max_file_mb} MB limit "
            "(raise MAX_FILE_MB to override)"
        )


def parse(path: Path) -> DocModel:
    """Dispatch a file to the right parser. Spreadsheets are handled upstream
    in the pipeline (they become SQL tables), not here."""
    guard(path)
    kind = _resolve_kind(path)
    if kind == "pdf":
        return _parse_pdf(path)
    if kind == "docx":
        return _parse_docx(path)
    if kind == "pptx":
        return _parse_pptx(path)
    if kind == "html":
        return _parse_html(path)
    if kind == "text":
        return _parse_text(path)
    if kind == "image":
        return _parse_image(path)
    if kind == "code":
        return _parse_code(path, CODE_EXT.get(path.suffix.lower(), "text"))
    raise ValueError(f"Unsupported file type for parse(): {path.name}")


def _doc_id(path: Path) -> str:
    return stable_id(path.name, str(path.stat().st_size))


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _parse_pdf(path: Path) -> DocModel:
    import pymupdf as fitz

    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ValueError(f"could not open PDF: {exc}") from exc

    if doc.needs_pass:
        doc.close()
        raise ValueError("PDF is password-protected")

    n_pages = min(doc.page_count, settings.max_pdf_pages)
    truncated = doc.page_count > n_pages
    doc_id = _doc_id(path)
    elements: list[Element] = []
    section_path: list[str] = []

    # First pass: body font size + which margin lines are running headers/footers.
    sizes: list[float] = []
    margin_lines: dict[str, set[int]] = {}
    for pi in range(n_pages):
        try:
            page = doc[pi]
            page_h = page.rect.height or 1.0
            for block in page.get_text("dict", sort=True).get("blocks", []):
                spans = [
                    s for line in block.get("lines", []) for s in line.get("spans", [])
                ]
                for span in spans:
                    if span.get("text", "").strip():
                        sizes.append(round(span["size"], 1))
                btext = clean_text("".join(s.get("text", "") for s in spans))
                y0, y1 = block.get("bbox", (0, 0, 0, 0))[1], block.get("bbox", (0, 0, 0, 0))[3]
                in_margin = y0 < 0.12 * page_h or y1 > 0.88 * page_h
                if btext and in_margin and len(btext) <= 90 and "\n" not in btext:
                    margin_lines.setdefault(_norm_furniture(btext), set()).add(pi)
        except Exception:  # pragma: no cover - skip an unreadable page
            continue
    body_size = statistics.median(sizes) if sizes else 10.0
    furniture = {
        norm for norm, pages in margin_lines.items()
        if len(pages) >= max(3, round(0.6 * n_pages))
    }

    total_text_chars = 0
    chart_budget = [_MAX_CHART_PAGES_PER_DOC]
    for pno in range(1, n_pages + 1):
        try:
            elements, total_text_chars, section_path = _pdf_page(
                fitz, doc[pno - 1], pno, body_size, elements, total_text_chars,
                section_path, furniture, doc_id, chart_budget,
            )
        except Exception as exc:  # pragma: no cover - one bad page won't stop ingest
            print(f"[ingest] {path.name} page {pno} skipped: {exc}")

    if truncated:
        elements.append(Element(
            kind="paragraph", page=n_pages,
            text=f"[Note: only the first {n_pages} of {doc.page_count} pages were "
                 f"ingested; raise MAX_PDF_PAGES to include the rest.]",
        ))

    # Scanned PDF: almost no extractable text -> OCR each page image.
    if total_text_chars < 40 * n_pages:
        elements.extend(_ocr_pdf_pages(doc, n_pages))

    title = _first_title(elements) or path.stem
    model = DocModel(
        doc_id=doc_id,
        source_file=path.name,
        source_kind=SourceKind.DOCUMENT,
        title=title,
        elements=elements,
        page_count=doc.page_count,
    )
    doc.close()
    return model


# "1 | P a g e", "Page 3", "Page 2 of 10", a bare "7" — with whitespace removed.
_PAGE_NUM_RE = re.compile(r"^(?:page)?\d{1,4}(?:\|?p\W?a\W?g\W?e)?(?:of\d{1,4})?$")


def _norm_furniture(text: str) -> str:
    """Normalise a margin line for repetition counting: fold case, drop spaces,
    and collapse digit runs so "Page 3" and "Page 4" count as the same line."""
    return re.sub(r"\d+", "#", re.sub(r"\s+", "", text.lower()))


def _is_page_furniture(text: str, furniture: set[str]) -> bool:
    flat = re.sub(r"\s+", "", text.lower())
    return bool(_PAGE_NUM_RE.match(flat)) or _norm_furniture(text) in furniture


def _pdf_page(
    fitz, page, pno: int, body_size: float,  # noqa: ANN001
    elements: list[Element], total_text_chars: int, section_path: list[str],
    furniture: set[str], doc_id: str, chart_budget: list[int],
) -> tuple[list[Element], int, list[str]]:
    """Append one page's elements. Returns updated (elements, char count, path)."""
    page_dict = page.get_text("dict", sort=True)
    # A large image on a page with essentially no other text is cover art / a
    # section-divider graphic / a watermark — not a data figure. Skip it
    # rather than spend an OCR + vision pass turning it into a junk chunk.
    page_is_blank = len(page.get_text("text").strip()) < 30

    # find_tables() is the slowest step per page (~0.12s each — minutes over a
    # long filing) and a pure-prose page has no table to find. Only run it
    # where the page actually has ruled lines: a real table draws a grid, and
    # a page with none is text. (A fully borderless table is missed, but the
    # finder's default strategy is unreliable on those anyway.)
    try:
        has_rules = len(page.get_drawings()) >= _MIN_TABLE_DRAWINGS
    except Exception:  # pragma: no cover
        has_rules = True
    try:
        tables = page.find_tables() if has_rules else []
    except Exception:  # pragma: no cover - finder is best-effort
        tables = []
    table_bboxes = []
    for tbl in getattr(tables, "tables", []):
        try:
            md = _table_to_markdown(tbl.extract())
        except Exception:  # pragma: no cover
            continue
        if md:
            table_bboxes.append(fitz.Rect(tbl.bbox))
            elements.append(Element(
                kind="table", text=md, table_markdown=md, page=pno,
                section_path=list(section_path), bbox=tuple(tbl.bbox),
            ))

    for block in page_dict.get("blocks", []):
        if block.get("type") == 1:  # image block
            # A chart/figure/diagram embedded in an otherwise-text page — the
            # whole-page OCR fallback below never fires here (this page has
            # plenty of extractable text), so without this, any data printed
            # *inside* the image (axis labels, data points, a puzzle's digits)
            # is invisible to the whole system. Worth a shot at whatever's
            # printed in it; a photo/logo just OCRs to nothing and is skipped.
            if page_is_blank:
                continue  # cover art / divider graphic — not worth reading
            el = _ocr_inline_image(block, pno, section_path, doc_id)
            if el:
                elements.append(el)
                total_text_chars += len(el.text)
            continue
        bbox = fitz.Rect(block["bbox"])
        if any(bbox.intersects(tb) for tb in table_bboxes):
            continue

        text_parts: list[str] = []
        max_size = body_size
        bold = False
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                t = span.get("text", "")
                if t.strip():
                    text_parts.append(t)
                    max_size = max(max_size, span.get("size", body_size))
                    if span.get("flags", 0) & 2 ** 4:
                        bold = True
            text_parts.append("\n")
        text = clean_text("".join(text_parts))
        if not text:
            continue
        if _is_page_furniture(text, furniture):
            continue  # running header/footer, page number — never structure
        total_text_chars += len(text)

        is_heading = (
            len(text) < 200
            and "\n" not in text
            and (max_size >= body_size * 1.12 or (bold and max_size >= body_size))
        )
        if is_heading:
            level = _heading_level(max_size, body_size)
            section_path = section_path[: level - 1] + [text]
            elements.append(Element(
                kind="heading", text=text, page=pno, heading_level=level,
                section_path=list(section_path), bbox=tuple(block["bbox"]),
            ))
        else:
            elements.append(Element(
                kind="paragraph", text=text, page=pno,
                section_path=list(section_path), bbox=tuple(block["bbox"]),
            ))

    fig = _vision_read_chart_page(fitz, page, pno, section_path, doc_id, chart_budget)
    if fig:
        elements.append(fig)

    return elements, total_text_chars, section_path


# A page with fewer drawn line segments than this is treated as prose and
# skipped by the (slow) table finder. All real tables in a ruled filing draw
# far more than this; pure text pages draw zero.
_MIN_TABLE_DRAWINGS = 6

# Skip tiny images (bullet icons, logos, decorative rules) — not worth an OCR
# pass and rarely carry data worth indexing.
_MIN_INLINE_IMAGE_PX = 80
# A small diagram (icons, dense number grids) OCRs unreliably at its native
# size, but there's no single scale that's best for every image — verified
# empirically that upscaling can as easily *introduce* a misread as fix one
# (e.g. a confident wrong digit at one scale, correct at another). Try a
# spread of scales and keep whichever pass actually found the most text,
# breaking ties by confidence — validated against 4 real embedded images
# (a missing row at native size, a misread digit at one scale) and it
# recovers the fully-correct reading in every case, where no fixed scale did.
# ONLY for genuinely small images: a 5x resize of a 2000px scan is 100+
# megapixels and takes ~6s each — minutes over a filing's exhibit images —
# for no benefit (a large image is already legible).
_OCR_SCALES = (1.0, 2.0, 3.0, 5.0)
_OCR_MULTISCALE_MAX_PX = 1000  # longest edge above which only native size is tried
_OCR_MAX_PX = 2600  # a bigger input just makes OCR slower — text detection is fine here


def _ocr_pil_image(img, w: float, h: float) -> str:  # noqa: ANN001
    """The multi-scale OCR sweep, factored out of ``_ocr_inline_image`` so it
    can run alongside (not instead of) a vision read: OCR is free and often
    good enough on its own, so it's kept as the always-on fallback for a
    provider/setup with no vision access. Returns ``""`` when OCR is
    unavailable or found nothing."""
    engine = _ocr_engine()
    if engine is None:
        return ""
    import numpy as np
    from PIL import Image

    # A huge embedded image (a filing's fold-out exhibit can be 9000px+) costs
    # seconds per OCR pass and per encode for no gain — shrink it first.
    if max(w, h) > _OCR_MAX_PX:
        s = _OCR_MAX_PX / max(w, h)
        img = img.resize((max(1, round(w * s)), max(1, round(h * s))), Image.LANCZOS)
        w, h = img.width, img.height

    scales = _OCR_SCALES if max(w, h) <= _OCR_MULTISCALE_MAX_PX else (1.0,)
    best: tuple[int, float, str] | None = None  # (n_detections, avg_confidence, text)
    for scale in scales:
        sized = (
            img.resize((round(w * scale), round(h * scale)), Image.LANCZOS)
            if scale != 1.0 else img
        )
        try:
            result, _elapsed = engine(np.ascontiguousarray(np.asarray(sized)))
        except Exception:  # pragma: no cover - one scale failing shouldn't sink the rest
            continue
        n = len(result) if result else 0
        if n == 0:
            continue
        avg_conf = sum(r[2] for r in result) / n
        if best is None or (n, avg_conf) > (best[0], best[1]):
            best = (n, avg_conf, clean_text("\n".join(r[1] for r in result)))
    return best[2] if best else ""


# A vision model reads a small, dense image (a cramped number grid, fine axis
# labels) poorly at its native size — many embedded figures are only 150-300px.
# Upscale the short edge to at least this before sending, well under the
# ~1568px where the API would downscale again. LANCZOS keeps printed digits
# crisp; a photo just gets slightly soft, which costs nothing here.
_VISION_MIN_PX = 900
# The Anthropic API rejects an image whose longer edge exceeds 8000px; stay
# comfortably under it. (It downscales anything over ~1568px internally
# anyway, so there's no quality loss from capping here.)
_VISION_MAX_PX = 7000


def _to_png_bytes(img, *, for_vision: bool = False) -> bytes:  # noqa: ANN001
    import io

    if for_vision:
        from PIL import Image

        short, long = min(img.size), max(img.size)
        scale = 1.0
        if 0 < short < _VISION_MIN_PX:
            scale = _VISION_MIN_PX / short
        if long * scale > _VISION_MAX_PX:  # cap wins over the upscale floor
            scale = _VISION_MAX_PX / long
        if scale != 1.0:
            img = img.resize(
                (max(1, round(img.width * scale)), max(1, round(img.height * scale))),
                Image.LANCZOS,
            )
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _save_image(doc_id: str, ref_id: str, png_bytes: bytes) -> str:
    """Persist one image to ``storage_dir/images/<doc_id>/<ref_id>.png`` and
    return the relative ref an ``Element``/``Chunk`` stores. Resolved back to
    a path with ``grounded.util.image_store_path`` — an agent (fallback,
    notably) reads these bytes at answer time for a fresh vision look at the
    real source image, not just its ingest-time text description."""
    from grounded.util import image_store_path

    path = image_store_path(f"{doc_id}/{ref_id}.png")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(png_bytes)
    return f"{doc_id}/{ref_id}.png"


def _describe_image_safe(png_bytes: bytes, prompt: str, pno: int) -> str:
    """Best-effort vision read: never lets a network/API hiccup sink ingest,
    same spirit as the try/except already wrapping OCR and page parsing."""
    from grounded.llm import chat_with_image

    try:
        return chat_with_image([{"role": "user", "content": prompt}], png_bytes).strip()
    except Exception as exc:  # pragma: no cover - network/API best-effort
        print(f"[ingest] vision read of page {pno} image failed: {exc}")
        return ""


_INLINE_IMAGE_PROMPT = (
    "Transcribe everything meaningful in this image as plain text — any "
    "numbers, letters, patterns, labels, or data. If it's a number/pattern "
    "puzzle, list the exact sequence shown, reading in the same order and "
    "orientation as printed. Be precise and literal; do not solve or "
    "interpret it, just transcribe what is printed."
)


def _ocr_inline_image(
    block: dict, pno: int, section_path: list[str], doc_id: str,
) -> Element | None:
    """Read one embedded image block from a text page's ``get_text("dict")``
    output — OCR always, plus a vision read when a vision-capable provider is
    configured (preferred when both succeed: OCR flattens a puzzle's 2D
    layout into a best-guess reading order, where a vision model can describe
    it as actually laid out). Returns ``None`` for anything too small or that
    yields nothing either way (a logo, a decorative photo) — cheap to try,
    safe to discard.

    Tagged ``paragraph``, not ``caption``: a "caption" is atomic in
    ``ingest/enumerate.py`` and would break a numbered run (e.g. Q1-30) in two
    if the image sits between two questions. As flowing text it just becomes
    part of whichever nearby item's chunk absorbs it — not perfectly
    attributed, but real data findable by search beats no data at all, and
    retrieval's relevance-floor padding (see ``agents/retrieval.py``) still
    surfaces it for nearby related questions even when it isn't the exact
    pinned chunk."""
    w, h = block.get("width", 0), block.get("height", 0)
    if w < _MIN_INLINE_IMAGE_PX or h < _MIN_INLINE_IMAGE_PX:
        return None
    raw = block.get("image")
    if not raw:
        return None
    try:
        import io

        from PIL import Image

        img = Image.open(io.BytesIO(raw)).convert("RGB")
    except Exception as exc:  # pragma: no cover - a malformed embedded image
        print(f"[ingest] could not decode embedded image on page {pno}: {exc}")
        return None

    ocr_text = _ocr_pil_image(img, w, h)
    vision_text = ""
    png_bytes = b""
    if vision_available():
        png_bytes = _to_png_bytes(img, for_vision=True)
        vision_text = _describe_image_safe(png_bytes, _INLINE_IMAGE_PROMPT, pno)

    text = vision_text or ocr_text
    if not text:
        return None
    ref = None
    if vision_text:
        ref_id = stable_id(doc_id, "img", str(pno), str(w), str(h), str(len(raw)))
        # Save the same upscaled bytes the vision call saw, so a later re-read
        # (the fallback agent) works from an image of the same quality.
        ref = _save_image(doc_id, ref_id, png_bytes or _to_png_bytes(img, for_vision=True))
    bbox = block.get("bbox")
    return Element(
        kind="paragraph", text=text, page=pno,
        section_path=list(section_path), image_ref=ref,
        bbox=tuple(bbox) if bbox else None,
    )


# A line/bar chart is very often drawn with vector paths (lines, ticks, axis
# labels as glyph runs) rather than embedded as a raster image at all —
# ``_ocr_inline_image`` above never sees it. Such a page gets rendered to a
# pixmap and read by vision. Gating it is delicate: a dense financial table
# (a 10-K has dozens) also draws a lot of ruled lines. Trigger on EITHER a
# recognisable bar-chart shape (see ``_has_bar_chart``) OR a page that's
# almost all graphic with little text — and cap the count per document.
_MIN_CHART_DRAWING_OPS = 40
_MAX_CHART_PAGE_TEXT = 400  # chars of real text above which it's a table, not a chart
_MAX_CHART_PAGES_PER_DOC = 4


def _has_bar_chart(page) -> bool:  # noqa: ANN001
    """True when the page draws a row of >=4 same-width, bottom-aligned,
    evenly-spaced filled rectangles — the geometry of a bar chart. This is
    what tells a real chart apart from a ruled table (whose fills are cell or
    header shading, not a row of bars). The 10-K's DAP and revenue charts
    match; none of its dozens of table pages do."""
    try:
        rects = [
            d["rect"] for d in page.get_drawings()
            if d.get("type") in ("f", "fs")
            and d["rect"].width > 2 and d["rect"].height > 2
        ]
    except Exception:  # pragma: no cover
        return False
    if len(rects) < 4:
        return False
    widths = sorted(r.width for r in rects)
    w_med = widths[len(widths) // 2]
    bars = [r for r in rects if w_med > 0 and abs(r.width - w_med) <= 0.35 * w_med]
    if len(bars) < 4:
        return False
    heights = [r.height for r in bars]
    if max(heights) - min(heights) <= 2:
        # identical-height boxes in a row -> a strip of labels/markers, not a
        # bar chart (bars encode their value in height, so real ones vary)
        return False
    bases = sorted(r.y1 for r in bars)
    if bases[-1] - bases[0] > 8:  # bottoms not aligned -> not one bar row
        return False
    xs = sorted(r.x0 for r in bars)
    gaps = [b - a for a, b in zip(xs, xs[1:], strict=False)]
    avg = sum(gaps) / len(gaps) if gaps else 0
    return bool(gaps) and avg > 0 and max(gaps) <= 3 * avg

_CHART_PAGE_PROMPT = (
    "This page contains a chart or graph. Reply with two parts:\n\n"
    "1. One sentence beginning 'This chart shows ...' that names, in plain "
    "words a reader would search for, what it plots: the full metric name "
    "with any abbreviation spelled out on first use (e.g. 'Family of Apps "
    "(FoA) revenue'), what the x-axis steps through and its range (e.g. 'by "
    "quarter from Q4 2022 to Q4 2024'), the units, and every series shown.\n\n"
    "2. Then the exact data as a markdown table: one row per x-axis category, "
    "one column per line/series, with the precise value at each point — read "
    "labels printed directly on the chart where present, and name each series "
    "as labeled. If a value truly can't be read confidently, write 'unclear' "
    "for that cell rather than guessing."
)


def _vision_read_chart_page(
    fitz, page, pno: int, section_path: list[str], doc_id: str,  # noqa: ANN001
    budget: list[int],
) -> Element | None:
    """Render a vector-drawn chart page to an image and describe it with
    vision — the one case OCR structurally cannot help with (see module
    docstring above ``_MIN_CHART_DRAWING_OPS``). ``None`` on any page that
    isn't chart-like, has no vision provider configured, is out of the
    per-document budget, or fails to render/describe — ingest must never
    break, hang, or run up a bill because a page looked chart-ish."""
    if not vision_available() or budget[0] <= 0:
        return None
    try:
        if _has_bar_chart(page):
            pass  # a real chart shape — read it even on a text-heavy page
        elif len(page.get_drawings()) < _MIN_CHART_DRAWING_OPS:
            return None
        elif len(page.get_text("text").strip()) > _MAX_CHART_PAGE_TEXT:
            # mostly drawn lines but full of real text -> a ruled table
            return None
    except Exception:  # pragma: no cover - best-effort
        return None
    budget[0] -= 1
    try:
        rect = page.rect
        # 2x for legibility, but clamp so the render can't exceed the API's
        # 8000px limit on an oversized page (fold-out exhibits, etc.).
        scale = min(2.0, _VISION_MAX_PX / max(rect.width, rect.height, 1.0))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale))
        png_bytes = pix.tobytes("png")
    except Exception as exc:  # pragma: no cover
        print(f"[ingest] could not render page {pno} for chart vision read: {exc}")
        return None
    text = _describe_image_safe(png_bytes, _CHART_PAGE_PROMPT, pno)
    if not text:
        return None
    ref = _save_image(doc_id, stable_id(doc_id, "chart", str(pno)), png_bytes)
    return Element(
        kind="figure", text=text, page=pno,
        section_path=list(section_path), image_ref=ref,
    )


def _ocr_pdf_pages(doc, n_pages: int) -> list[Element]:  # noqa: ANN001
    import numpy as np
    import pymupdf as fitz

    engine = _ocr_engine()
    if engine is None:
        return []
    out: list[Element] = []
    for pno in range(1, min(n_pages, 50) + 1):  # OCR is slow; cap it
        try:
            pix = doc[pno - 1].get_pixmap(matrix=fitz.Matrix(2, 2))
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
                pix.height, pix.width, pix.n
            )
            if pix.n >= 4:
                img = img[:, :, :3]
            text = _run_ocr(engine, np.ascontiguousarray(img))
        except Exception as exc:  # pragma: no cover
            print(f"[ingest] OCR of page {pno} failed: {exc}")
            continue
        if text:
            out.append(Element(kind="paragraph", text=text, page=pno))
    return out


def _table_to_markdown(rows: list[list], *, max_cols: int = 40, max_rows: int = 400,
                       max_cell: int = 300) -> str:
    def cell(v: object) -> str:
        s = ("" if v is None else str(v)).strip().replace("\n", " ").replace("|", "\\|")
        return s if len(s) <= max_cell else s[:max_cell] + "…"

    rows = [[cell(c) for c in r[:max_cols]] for r in rows]
    rows = [r for r in rows if any(r)]
    if len(rows) < 2:
        return ""
    truncated_rows = len(rows) > max_rows
    rows = rows[:max_rows]
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    if truncated_rows:
        lines.append(f"| _(table truncated at {max_rows} rows)_ |")
    return "\n".join(lines)


def _heading_level(size: float, body: float) -> int:
    ratio = size / body
    if ratio >= 1.6:
        return 1
    if ratio >= 1.35:
        return 2
    if ratio >= 1.18:
        return 3
    return 4


def _first_title(elements: list[Element]) -> str | None:
    for el in elements:
        if el.kind in {"title", "heading"}:
            return el.text
    return None


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _parse_docx(path: Path) -> DocModel:
    import docx
    from docx.document import Document as _Doc
    from docx.table import Table as _Table
    from docx.text.paragraph import Paragraph as _Para

    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise ValueError(f"could not open DOCX: {exc}") from exc

    elements: list[Element] = []
    section_path: list[str] = []

    def iter_body(parent):  # yields paragraphs and tables in document order
        node = parent.element.body if isinstance(parent, _Doc) else parent._tc
        for child in node.iterchildren():
            if child.tag.endswith("}p"):
                yield _Para(child, parent)
            elif child.tag.endswith("}tbl"):
                yield _Table(child, parent)

    for item in iter_body(d):
        if isinstance(item, _Table):
            rows = [[c.text for c in row.cells] for row in item.rows]
            md = _table_to_markdown(rows)
            if md:
                elements.append(Element(
                    kind="table", text=md, table_markdown=md,
                    section_path=list(section_path),
                ))
            continue

        text = clean_text(item.text)
        if not text:
            continue
        style = (item.style.name or "").lower() if item.style else ""
        if style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = int(digits) if digits else 1
            section_path = section_path[: level - 1] + [text]
            elements.append(Element(
                kind="heading", text=text, heading_level=level,
                section_path=list(section_path),
            ))
        elif style == "title":
            elements.append(Element(kind="title", text=text, heading_level=1))
        else:
            kind = "list_item" if "list" in style else "paragraph"
            elements.append(Element(kind=kind, text=text, section_path=list(section_path)))

    if not elements:
        elements.append(Element(kind="paragraph", text=f"[DOCX file {path.name} "
                                "contained no extractable text — it may be images only.]"))

    return DocModel(
        doc_id=_doc_id(path),
        source_file=path.name,
        source_kind=SourceKind.DOCUMENT,
        title=_first_title(elements) or path.stem,
        elements=elements,
    )


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def _pptx_chart_element(chart, page: int, title: str) -> Element | None:  # noqa: ANN001
    """Turn a native PowerPoint chart into a markdown table plus a one-line
    caption. The chart's data lives in its own XML part (categories + one or
    more named series), so this is exact — no rendering or vision needed. A
    malformed or empty chart yields ``None`` rather than raising."""
    try:
        cats = [str(c) for c in chart.plots[0].categories]
        series = [
            (s.name or f"Series {i + 1}", list(s.values))
            for i, s in enumerate(chart.series)
        ]
    except Exception:  # pragma: no cover - defensive against odd chart XML
        return None
    if not cats or not series or not any(vals for _, vals in series):
        return None
    try:
        kind = chart.chart_type.name.replace("_", " ").lower()
    except Exception:  # pragma: no cover
        kind = "chart"

    def _fmt(v: object) -> str:
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        if isinstance(v, int):
            return f"{v:,}"
        if isinstance(v, float):
            return f"{v:.2f}".rstrip("0").rstrip(".")
        return str(v)

    header = "| Category | " + " | ".join(name for name, _ in series) + " |"
    sep = "| --- |" + " --- |" * len(series)
    body = [
        f"| {cat} | "
        + " | ".join(_fmt(vals[r] if r < len(vals) else None) for _, vals in series)
        + " |"
        for r, cat in enumerate(cats)
    ]
    md = "\n".join([header, sep, *body])
    names = ", ".join(name for name, _ in series)
    caption = (
        f"{kind} chart{' — ' + title if title else ''}: "
        f"{names} by {cats[0]}…{cats[-1]}."
    )
    return Element(
        kind="table", text=f"{caption}\n\n{md}", table_markdown=md,
        page=page, section_path=[title] if title else [],
    )


def _parse_pptx(path: Path) -> DocModel:
    from pptx import Presentation

    prs = Presentation(str(path))
    elements: list[Element] = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title and slide.shapes.title.text.strip():
            title = clean_text(slide.shapes.title.text)
            elements.append(
                Element(kind="heading", text=title, page=idx, heading_level=1,
                        section_path=[title])
            )
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_table:
                rows = [
                    [cell.text for cell in row.cells]
                    for row in shape.table.rows
                ]
                md = _table_to_markdown(rows)
                if md:
                    elements.append(
                        Element(kind="table", text=md, table_markdown=md,
                                page=idx, section_path=[title] if title else [])
                    )
            elif getattr(shape, "has_chart", False):
                el = _pptx_chart_element(shape.chart, idx, title)
                if el:
                    elements.append(el)
            elif shape.has_text_frame:
                text = clean_text(shape.text_frame.text)
                if text:
                    elements.append(
                        Element(kind="paragraph", text=text, page=idx,
                                section_path=[title] if title else [])
                    )
    return DocModel(
        doc_id=_doc_id(path),
        source_file=path.name,
        source_kind=SourceKind.DOCUMENT,
        title=_first_title(elements) or path.stem,
        elements=elements,
        page_count=len(prs.slides),
    )


# ---------------------------------------------------------------------------
# text / markdown
# ---------------------------------------------------------------------------


# "3. On-call", "3.1 Setup", "APPENDIX A" — heading shapes common in plain .txt
_TXT_HEADING = re.compile(r"^\s*(?:\d+(?:\.\d+)*\.?\s+\S|[A-Z][A-Z0-9 /&'-]{3,}$)")
_UNDERLINE = re.compile(r"^[=\-~*_]{3,}\s*$")


def _parse_html(path: Path) -> DocModel:
    """HTML in, Markdown-shaped elements out: convert to Markdown and reuse the
    Markdown path so headings become sections and tables/lists survive. Full
    web pages carry nav/footer noise; clean exported documents parse well."""
    import html2text

    conv = html2text.HTML2Text()
    conv.body_width = 0  # no hard-wrapping — the chunker handles layout
    conv.ignore_images = True
    conv.ignore_emphasis = True
    conv.pad_tables = True  # emit real "| a | b |" Markdown tables
    md = conv.handle(path.read_text(encoding="utf-8", errors="replace"))
    return _parse_text(path, text=md, is_md=True)


def _parse_text(
    path: Path, *, text: str | None = None, is_md: bool | None = None
) -> DocModel:
    raw = path.read_text(encoding="utf-8", errors="replace") if text is None else text
    if is_md is None:
        is_md = path.suffix.lower() in {".md", ".markdown"}
    lines_in = raw.splitlines()
    elements: list[Element] = []
    section_path: list[str] = []
    buf: list[str] = []

    def flush() -> None:
        if buf:
            text = clean_text("\n".join(buf))
            if text:
                elements.append(
                    Element(kind="paragraph", text=text,
                            section_path=list(section_path))
                )
            buf.clear()

    def add_heading(text: str, level: int) -> None:
        nonlocal section_path
        flush()
        section_path = section_path[: level - 1] + [text]
        elements.append(
            Element(kind="heading", text=text, heading_level=level,
                    section_path=list(section_path))
        )

    in_fence = False
    for i, line in enumerate(lines_in):
        stripped = line.strip()
        if stripped.startswith("```"):
            in_fence = not in_fence
            buf.append(line)
            continue
        if in_fence:
            buf.append(line)
            continue

        if line.lstrip().startswith("#"):
            hashes = len(line) - len(line.lstrip("#"))
            add_heading(line.lstrip("#").strip(), min(hashes, 6))
        elif _UNDERLINE.match(line) and buf and buf[-1].strip():
            title = buf.pop().strip()  # "Setext" heading: text then ==== / ----
            add_heading(title, 1)
        elif not stripped:
            flush()  # blank line ends a paragraph
        elif not is_md and _TXT_HEADING.match(line) and len(stripped) < 70 and (
            i + 1 < len(lines_in) and not lines_in[i + 1].strip()
        ):
            add_heading(stripped, 2 if stripped[0].isdigit() else 1)
        else:
            buf.append(line)
    flush()

    return DocModel(
        doc_id=_doc_id(path),
        source_file=path.name,
        source_kind=SourceKind.DOCUMENT,
        title=_first_title(elements) or path.stem,
        elements=elements,
    )


# ---------------------------------------------------------------------------
# images
# ---------------------------------------------------------------------------


def _prepare_image(path: Path):  # noqa: ANN202 - returns an HxWx3 uint8 ndarray
    """Validate, fix EXIF orientation, and downscale very large images for OCR."""
    import numpy as np
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        img = Image.open(path)
        img.load()
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError(f"unreadable image (format not supported): {exc}") from exc

    img = ImageOps.exif_transpose(img).convert("RGB")
    longest = max(img.size)
    if longest > settings.max_image_px:
        scale = settings.max_image_px / longest
        img = img.resize((int(img.width * scale), int(img.height * scale)))
    return np.ascontiguousarray(np.asarray(img))


def _parse_image(path: Path) -> DocModel:
    elements: list[Element] = []
    prepared = _prepare_image(path)

    engine = _ocr_engine()
    if engine is not None:
        text = _run_ocr(engine, prepared)
        if text:
            elements.append(Element(kind="paragraph", text=text, page=1))

    caption = _vlm_caption(path)
    if caption:
        elements.append(Element(kind="caption", text=caption, page=1))

    if not elements:
        elements.append(
            Element(kind="caption", text=f"Image file: {path.name} (no text detected).",
                    page=1)
        )

    return DocModel(
        doc_id=_doc_id(path),
        source_file=path.name,
        source_kind=SourceKind.IMAGE,
        title=path.stem,
        elements=elements,
        page_count=1,
    )


@functools.lru_cache(maxsize=1)
def _ocr_engine():  # noqa: ANN202
    try:
        from rapidocr_onnxruntime import RapidOCR

        return RapidOCR()
    except Exception as exc:  # pragma: no cover
        print(f"[ingest] OCR unavailable: {exc}")
        return None


def _run_ocr(engine, image) -> str:  # noqa: ANN001
    try:
        result, _ = engine(image)
    except Exception as exc:  # pragma: no cover
        print(f"[ingest] OCR failed: {exc}")
        return ""
    if not result:
        return ""
    return clean_text("\n".join(line[1] for line in result))


def _vlm_caption(path: Path) -> str:
    if not settings.vlm_model:
        return ""
    try:
        import ollama

        client = ollama.Client(host=settings.ollama_host)
        resp = client.chat(
            model=settings.vlm_model,
            messages=[
                {
                    "role": "user",
                    "content": (
                        "Describe this image in 2-4 sentences. If it is a chart, "
                        "state the chart type, axes, and the main takeaway. If it "
                        "contains a table, summarize its columns."
                    ),
                    "images": [str(path)],
                }
            ],
            options={"temperature": 0.1},
        )
        return clean_text(resp["message"]["content"])
    except Exception as exc:  # pragma: no cover
        print(f"[ingest] VLM caption failed: {exc}")
        return ""


# ---------------------------------------------------------------------------
# code
# ---------------------------------------------------------------------------

_DEF_RE = re.compile(
    r"^[ \t]*(?:def |class |func |function |public |private |protected |fn |"
    r"async def |export function |export default function )",
)


def _parse_code(path: Path, language: str) -> DocModel:
    raw = path.read_text(encoding="utf-8", errors="replace")
    lines = raw.splitlines()
    elements: list[Element] = []

    # Split on top-level definition boundaries; keep leading module code together.
    starts = [0] + [i for i, ln in enumerate(lines) if _DEF_RE.match(ln) and not ln[:1].isspace()]
    starts = sorted(set(starts))
    bounds = list(zip(starts, starts[1:] + [len(lines)], strict=True))
    for a, b in bounds:
        segment = "\n".join(lines[a:b]).strip("\n")
        if segment.strip():
            first = next((ln for ln in lines[a:b] if ln.strip()), "")
            elements.append(
                Element(
                    kind="code",
                    text=f"```{language}\n{segment}\n```",
                    section_path=[first.strip()[:80]] if _DEF_RE.match(first) else [],
                )
            )

    return DocModel(
        doc_id=_doc_id(path),
        source_file=path.name,
        source_kind=SourceKind.CODE,
        title=path.name,
        elements=elements,
        extra={"language": language, "lines": str(len(lines))},
    )
